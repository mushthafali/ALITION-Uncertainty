from __future__ import annotations

import io
import os
import sys
import json
from typing import Any, Dict, List, Optional

import matplotlib  # type: ignore
matplotlib.use('Agg')
import matplotlib.pyplot as plt  # type: ignore
import openpyxl  # type: ignore
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore
from reportlab.lib import colors  # type: ignore
from reportlab.lib.colors import Color  # type: ignore
from reportlab.lib.pagesizes import letter  # type: ignore
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # type: ignore
from reportlab.pdfgen import canvas  # type: ignore
from reportlab.platypus import (  # type: ignore
    Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak, HRFlowable
)

from ..models import CalibrationInput, CalibrationResult
from ..visualization import CalibrationCharts
from .storage import _model_dump

class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas that dynamically calculates the total page count and prints
    consistent running page numbers and certificate tracking on every page.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_footer(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_footer(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(Color(100/255, 116/255, 139/255))
        self.setStrokeColor(Color(226/255, 232/255, 240/255))
        self.setLineWidth(0.6)
        from reportlab.lib.units import mm
        self.line(16*mm, 12*mm, (210-16)*mm, 12*mm)
        cert_no_str = getattr(self, '_cert_number', '')
        left_text = f"Sertifikat / Certificate: {cert_no_str}" if cert_no_str else "Alition Calibration System"
        self.drawString(16*mm, 7.5*mm, left_text)
        page_text = f"Halaman {self._pageNumber} dari {page_count}  |  Page {self._pageNumber} of {page_count}"
        self.drawRightString((210-16)*mm, 7.5*mm, page_text)
        self.restoreState()


def _get_logo_path() -> Optional[str]:
    candidates = [
        os.path.join(getattr(sys, '_MEIPASS', ''), 'img', 'logo_polos.png'),
        os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..', 'img', 'logo_polos.png')),
        os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', 'img', 'logo_polos.png')),
        os.path.join(os.getcwd(), 'img', 'logo_polos.png'),
    ]
    for p in candidates:
        if p and os.path.exists(p):
            return p
    return None

def make_numbered_canvas(cert_number: str):
    class DocNumberedCanvas(NumberedCanvas):
        _cert_number = cert_number
    return DocNumberedCanvas

def _format_id_date(dt_str: str) -> str:
    """Format ISO date string to Indonesian format (e.g. 07 Juli 2026) safely."""
    import datetime
    try:
        dt = datetime.datetime.strptime(dt_str[:10], "%Y-%m-%d")
    except Exception:
        dt = datetime.datetime.now()
    months = [
        "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember"
    ]
    return f"{dt.day:02d} {months[dt.month - 1]} {dt.year}"
class ReportingModule:
    def __init__(self, input_data: CalibrationInput, result: CalibrationResult):
        self.input_data = input_data
        self.result = result

    # ── Excel Export (dengan rumus komprehensif) ──────────────────────────
    def export_excel(self, filename: str):
        wb = openpyxl.Workbook()

        ws_sum = wb.active
        if ws_sum is None:
            ws_sum = wb.create_sheet("Summary")
        else:
            ws_sum.title = "Summary"
        ws_pts = wb.create_sheet("Point Data")
        ws_raw = wb.create_sheet("Raw Readings")

        hdr_fill  = PatternFill("solid", fgColor="1E2A3A")
        hdr_font  = Font(bold=True, color="FFFFFF", name="Calibri")
        sub_fill  = PatternFill("solid", fgColor="2C3E50")
        sub_font  = Font(bold=True, color="A0B4C8", name="Calibri")
        val_font  = Font(name="Calibri")
        thin_border = Border(
            left=Side(style='thin', color="C0C8D8"), right=Side(style='thin', color="C0C8D8"),
            top=Side(style='thin', color="C0C8D8"), bottom=Side(style='thin', color="C0C8D8")
        )

        def hdr(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.font = hdr_font; c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center"); c.border = thin_border
            return c

        def sub(ws, row, col, text):
            c = ws.cell(row=row, column=col, value=text)
            c.font = sub_font; c.fill = sub_fill
            c.alignment = Alignment(horizontal="left", vertical="center"); c.border = thin_border
            return c

        def val(ws, row, col, value, num_fmt=None):
            c = ws.cell(row=row, column=col, value=value)
            c.font = val_font; c.border = thin_border
            c.alignment = Alignment(horizontal="center")
            if num_fmt: c.number_format = num_fmt
            return c

        # ── 1. Populate Raw Readings ──
        hdr(ws_raw, 1, 1, "Set Point")
        max_reads = max((len(pr.readings) for pr in self.result.point_results), default=1)
        for j in range(1, max_reads + 1):
            hdr(ws_raw, 1, j + 1, f"Reading {j}")

        n_pts = len(self.result.point_results)
        for i, pr in enumerate(self.result.point_results, start=2):
            val(ws_raw, i, 1, pr.set_point, '0.00')
            for j, rd in enumerate(pr.readings, start=1):
                val(ws_raw, i, j + 1, rd, '0.0000')

        ws_raw.column_dimensions["A"].width = 12
        for j in range(2, max_reads + 2):
            ws_raw.column_dimensions[get_column_letter(j)].width = 14

        # ── 2. Populate Summary Metadata ──
        meta = self.input_data.metadata
        hdr(ws_sum, 1, 1, "ALITION - Calibration Report")
        ws_sum.merge_cells("A1:D1")
        ws_sum["A1"].font = Font(bold=True, size=14, color="FFFFFF", name="Calibri")

        info = [
            ("Operator", meta.operator_name if meta else ""),
            ("Instrument Model", meta.instrument_model if meta else ""),
            ("Date", meta.date if meta else ""),
            ("Location", meta.location if meta else ""),
            ("LRV", self.input_data.lrv),
            ("URV", self.input_data.urv),
            ("Tolerance (%)", self.input_data.tolerance_pct),
        ]
        for i, (label, v) in enumerate(info, start=2):
            sub(ws_sum, i, 1, label)
            if isinstance(v, float): val(ws_sum, i, 2, v, '0.0000')
            else: val(ws_sum, i, 2, v)

        extra_meta = [
            ("Tag Number", getattr(meta, "tag_number", "") if meta else ""),
            ("Serial Number", getattr(meta, "serial_number", "") if meta else ""),
        ]
        for i, (label, v) in enumerate(extra_meta, start=2):
            sub(ws_sum, i, 3, label)
            val(ws_sum, i, 4, v)
        ws_sum.column_dimensions["C"].width = 18
        ws_sum.column_dimensions["D"].width = 22

        # ── 3. Populate Point Data (Formulas) ──
        headers = ["Set Point", "Avg Reading", "Error (abs)", "Error (%FS)", "Std Dev", "Repeatability Err", "Repeatability %FS", "Status", "Linearity Residual"]
        for col, h in enumerate(headers, 1):
            hdr(ws_pts, 1, col, h)

        last_col_let = get_column_letter(max_reads + 1)
        
        for i in range(2, n_pts + 2):
            # A: Set Point
            val(ws_pts, i, 1, f"='Raw Readings'!A{i}", '0.00')
            # B: Avg Reading
            val(ws_pts, i, 2, f"=IF(ISERROR(AVERAGE('Raw Readings'!B{i}:{last_col_let}{i})), 0, AVERAGE('Raw Readings'!B{i}:{last_col_let}{i}))", '0.0000')
            # C: Error
            val(ws_pts, i, 3, f"=B{i}-A{i}", '0.00000')
            # D: Error %FS
            val(ws_pts, i, 4, f"=C{i}/Summary!$B$10*100", '0.0000"%"')
            # E: Std Dev
            val(ws_pts, i, 5, f"=IF(ISERROR(STDEV('Raw Readings'!B{i}:{last_col_let}{i})), 0, STDEV('Raw Readings'!B{i}:{last_col_let}{i}))", '0.00000')
            
            # F: Repeatability Err = MAX(ABS(Readings - Avg))
            abs_diffs = []
            for j in range(2, max_reads + 2):
                c_let = get_column_letter(j)
                abs_diffs.append(f"ABS('Raw Readings'!{c_let}{i}-B{i})")
            rep_err_formula = f"=MAX({','.join(abs_diffs)})"
            val(ws_pts, i, 6, rep_err_formula, '0.00000')
            
            # G: Repeatability %FS
            val(ws_pts, i, 7, f"=F{i}/Summary!$B$10*100", '0.0000"%"')
            # H: Status
            val(ws_pts, i, 8, f'=IF(ABS(D{i})<=Summary!$B$8, "PASS", "FAIL")')
            ws_pts.cell(row=i, column=8).font = Font(bold=True, name="Calibri")
            # I: Linearity Residual
            val(ws_pts, i, 9, f"=B{i} - (Summary!$B$16 * A{i} + Summary!$B$17)", '0.00000')

        for ci, w in enumerate([12, 14, 14, 14, 12, 18, 18, 10, 18], 1):
            ws_pts.column_dimensions[get_column_letter(ci)].width = w

        # ── 4. Populate Summary Metrics (Formulas) ──
        row_start = 9
        hdr(ws_sum, row_start, 1, "Metric")
        hdr(ws_sum, row_start, 2, "Value")
        
        metrics = [
            ("Span", f"=B7-B6"), # B10
            ("Max Error (%FS)", f"=MAX(ABS(MIN('Point Data'!D2:D{n_pts+1})), ABS(MAX('Point Data'!D2:D{n_pts+1})))"), # B11
            ("Actual Accuracy (%)", f"=B11"), # B12
            ("Mean Bias Error", f"=AVERAGE('Point Data'!C2:C{n_pts+1})"), # B13
            ("Std Deviation", f"=IF(ISERROR(STDEV('Point Data'!C2:C{n_pts+1})), 0, STDEV('Point Data'!C2:C{n_pts+1}))"), # B14
            ("RMSE", f"=SQRT(SUMSQ('Point Data'!C2:C{n_pts+1})/COUNT('Point Data'!C2:C{n_pts+1}))"), # B15
            ("Slope (Linearity)", f"=SLOPE('Point Data'!B2:B{n_pts+1}, 'Point Data'!A2:A{n_pts+1})"), # B16
            ("Intercept (Linearity)", f"=INTERCEPT('Point Data'!B2:B{n_pts+1}, 'Point Data'!A2:A{n_pts+1})"), # B17
            ("Linearity Error (%FS)", f"=MAX(ABS(MIN('Point Data'!I2:I{n_pts+1})), ABS(MAX('Point Data'!I2:I{n_pts+1}))) / B10 * 100"), # B18
            ("Zero Error", f"='Point Data'!C2"), # B19
            ("Span Error", f"='Point Data'!C{n_pts+1} - 'Point Data'!C2"), # B20
            ("Overall Status", f'=IF(B11<=B8, "PASS", "FAIL")'), # B21
        ]
        
        for i, (label, formula) in enumerate(metrics, start=row_start + 1):
            sub(ws_sum, i, 1, label)
            c = val(ws_sum, i, 2, formula, '0.00000')
            if label == "Overall Status":
                c.number_format = "@"
                c.font = Font(bold=True, name="Calibri")
            elif "(%FS)" in label or "(%)" in label:
                if label != "Actual Accuracy (%)":
                    # For format consistency, keep as number or add % sign
                    pass

        ws_sum.column_dimensions["A"].width = 26
        ws_sum.column_dimensions["B"].width = 22

        wb.save(filename)

    # ── CSV (opsional, tetap tersedia) ─────────────────────────────────────
    def export_csv(self, filename: str):
        import pandas as pd  # type: ignore
        data = []
        for pr in self.result.point_results:
            data.append({
                "Set Point": pr.set_point,
                "Average Reading": pr.avg_reading,
                "Error": pr.error,
                "% Span Error": pr.error_pct_span,
                "Std Dev": pr.std_dev,
                "Repeatability Error": pr.repeatability_error,
                "Repeatability % Span": pr.repeatability_pct_span,
                "Status": pr.status,
            })
        pd.DataFrame(data).to_csv(filename, index=False)

    def export_json(self, filename: str) -> None:
        data = {"input": _model_dump(self.input_data), "result": _model_dump(self.result)}
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=4)

    # ── PDF Export ──────────────────────────────────────────────────────────
    def export_pdf(self, filename: str, cert_number: str = ''):
        import io
        import os
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.units import mm  # type: ignore
        from .storage import StorageManager

        if not cert_number:
            meta_date = self.input_data.metadata.date if self.input_data.metadata else None
            cert_number = StorageManager.generate_certificate_number(meta_date)
        meta_date = self.input_data.metadata.date if self.input_data.metadata else None
        import datetime
        export_date = str(meta_date or datetime.datetime.now().strftime('%Y-%m-%d'))

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            leftMargin=16*mm, rightMargin=16*mm,
            topMargin=14*mm, bottomMargin=14*mm
        )
        styles = getSampleStyleSheet()
        elements: List[Any] = []

        DARK   = Color(30/255, 41/255, 59/255)
        GREY_C = Color(100/255, 116/255, 139/255)
        LABEL  = Color(51/255, 65/255, 85/255)
        ACCENT = Color(30/255, 64/255, 175/255)
        LINE_C = Color(203/255, 213/255, 225/255)

        # ── Plain White Header with Alition Identity ──────────────────────────
        h_brand_title = Paragraph("<b>ALITION</b>", ParagraphStyle('HBrandTitle', fontName='Helvetica-Bold', fontSize=16, leading=19, textColor=DARK))
        h_brand_sub = Paragraph("AUTOMATIC CALIBRATION CALCULATIONS SYSTEM", ParagraphStyle('HBrandSub', fontName='Helvetica', fontSize=7, leading=9, textColor=GREY_C))

        logo_path = _get_logo_path()
        if logo_path and os.path.exists(logo_path):
            logo_flowable = Image(logo_path, width=34, height=34)
            header_table = Table([
                [logo_flowable, [h_brand_title, h_brand_sub]]
            ], colWidths=[42, 462])
            header_table.setStyle(TableStyle([
                ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING',   (0, 0), (0, 0), 0),
                ('RIGHTPADDING',  (0, 0), (0, 0), 0),
                ('LEFTPADDING',   (1, 0), (1, 0), 6),
                ('RIGHTPADDING',  (1, 0), (1, 0), 0),
                ('TOPPADDING',    (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('LINEBELOW',     (0, 0), (-1, -1), 1.2, ACCENT),
            ]))
        else:
            header_table = Table([
                [[h_brand_title, h_brand_sub]]
            ], colWidths=[504])
            header_table.setStyle(TableStyle([
                ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING',   (0, 0), (-1, -1), 0),
                ('RIGHTPADDING',  (0, 0), (-1, -1), 0),
                ('TOPPADDING',    (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('LINEBELOW',     (0, 0), (-1, -1), 1.2, ACCENT),
            ]))
        elements.append(header_table)
        elements.append(Spacer(1, 12))

        # ── Prominent Main Title (Highlight Utama - Formal Authoritative Serif) ─
        title_main = Paragraph("SERTIFIKAT KALIBRASI", ParagraphStyle(
            'MainTitle', fontName='Times-Bold', fontSize=19, leading=23, textColor=DARK, alignment=1, spaceAfter=2
        ))
        title_sub = Paragraph("CALIBRATION CERTIFICATE", ParagraphStyle(
            'MainSub', fontName='Times-Italic', fontSize=9.5, leading=12, textColor=GREY_C, alignment=1, spaceAfter=4
        ))
        # Monospaced tabular font for the official registry certificate number
        title_cert_no = Paragraph(f"<font name='Helvetica' size=8.5 color='#64748b'>Nomor Sertifikat / Certificate No : </font><font name='Courier-Bold' size=10.5 color='#1d4ed8'>{cert_number}</font>", ParagraphStyle(
            'MainCertNo', alignment=1
        ))

        title_block = Table([[title_main], [title_sub], [title_cert_no]], colWidths=[504])
        title_block.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        elements.append(title_block)
        elements.append(Spacer(1, 14))

        # ── Purposeful Typographic Hierarchy (Antislop UI) ─────────────────────
        # Section Headings
        s1_section = ParagraphStyle('S1Section', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=10, leading=13,
            textColor=DARK, spaceBefore=0, spaceAfter=1)
        s1_sub = ParagraphStyle('S1Sub', parent=styles['Normal'],
            fontName='Times-Italic', fontSize=8, leading=10,
            textColor=GREY_C, spaceAfter=4)

        # Labels & Colons
        s1_label = ParagraphStyle('S1Label', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=8, leading=9.5,
            textColor=LABEL)
        s1_colon = ParagraphStyle('S1Colon', parent=styles['Normal'],
            fontName='Helvetica', fontSize=8, leading=9.5,
            textColor=GREY_C)

        # Values categorized by functional role
        s1_val_text = ParagraphStyle('S1ValTxt', parent=styles['Normal'],
            fontName='Helvetica', fontSize=8, leading=9.5,
            textColor=DARK)
        s1_val_code = ParagraphStyle('S1ValCode', parent=styles['Normal'],
            fontName='Courier-Bold', fontSize=8, leading=9.5,
            textColor=DARK)
        s1_val_num  = ParagraphStyle('S1ValNum', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=8, leading=9.5,
            textColor=DARK)

        # Helper for bilingual labels (Bahasa Indonesia on top, English in italics below)
        def _lbl(id_text: str, en_text: str) -> Paragraph:
            return Paragraph(f"{id_text}<br/><font name='Times-Italic' size=6.8 color='#64748b'>{en_text}</font>", s1_label)

        # Formal Approval Sign-off (Legal authority serif)
        export_dt_formatted = _format_id_date(export_date)
        sig_title = Paragraph("<b>Disetujui oleh :</b><br/><font name='Times-Italic' size=8 color='#64748b'>Approved by :</font>", ParagraphStyle('SigTitle', fontName='Times-Bold', fontSize=9.5, leading=12, textColor=DARK))
        sig_loc_date = Paragraph(f"Jakarta, {export_dt_formatted}", ParagraphStyle('SigLocDate', fontName='Times-Roman', fontSize=9.5, leading=13, textColor=DARK))

        # ── Helper: thin horizontal rule ──
        def _hr():
            return HRFlowable(width='100%', thickness=0.5, color=LINE_C,
                              spaceBefore=2, spaceAfter=4)

        # ── Table style (reused for all detail tables) ──
        detail_ts = TableStyle([
            ('VALIGN',        (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('TOPPADDING',    (0, 0), (-1, -1), 1),
            ('LEFTPADDING',   (0, 0), (-1, -1), 0),
            ('RIGHTPADDING',  (0, 0), (-1, -1), 0),
        ])
        L_COL = [105, 8, 131]
        R_COL = [105, 8, 131]

        if self.input_data.metadata:
            meta = self.input_data.metadata
            elems = elements

            # ═══════════════════════════════════════════════════════════════════
            #  COLUMN 1 (LEFT): Identitas Alat & Data Kalibrasi
            # ═══════════════════════════════════════════════════════════════════
            left: List[Any] = []

            # ── 1. Identitas Alat (tag number tidak diikutsertakan) ──
            left.append(Paragraph("Identitas Alat", s1_section))
            left.append(Paragraph("Instrument Details", s1_sub))
            rows_alat = [
                [_lbl("Tipe Alat", "Instrument Type"),                     Paragraph(":", s1_colon), Paragraph("Pressure Gauge",                  s1_val_text)],
                [_lbl("Model", "Model"),                                   Paragraph(":", s1_colon), Paragraph(meta.instrument_model,            s1_val_text)],
                [_lbl("Nomor Seri", "Serial Number"),                     Paragraph(":", s1_colon), Paragraph(getattr(meta, 'serial_number', '') or '-', s1_val_code)],
                [_lbl("Rentang Ukur", "Measuring Range"),                 Paragraph(":", s1_colon), Paragraph("—",                               s1_val_num)],
                [_lbl("Toleransi / Resolusi", "Tolerance / Resolution"),  Paragraph(":", s1_colon), Paragraph("—",                               s1_val_num)],
                [_lbl("Lokasi", "Location"),                               Paragraph(":", s1_colon), Paragraph(meta.location,                     s1_val_text)],
            ]
            t = Table(rows_alat, colWidths=L_COL); t.setStyle(detail_ts)
            left.append(t)
            left.append(_hr())

            # ── 2. Data Kalibrasi ──
            left.append(Paragraph("Data Kalibrasi", s1_section))
            left.append(Paragraph("Calibration Data", s1_sub))
            rows_cal = [
                [_lbl("Tanggal Kalibrasi", "Calibration Date"),           Paragraph(":", s1_colon), Paragraph(meta.date,                            s1_val_num)],
                [_lbl("Tipe Laporan", "Report Type"),                     Paragraph(":", s1_colon), Paragraph("Calibration Analysis Report",         s1_val_text)],
                [_lbl("Kalibrator Standar", "Standard Calibrator"),       Paragraph(":", s1_colon), Paragraph("—",                                   s1_val_text)],
                [_lbl("Batas Kesalahan Izin (MPE)", "Max Permissible Error"), Paragraph(":", s1_colon), Paragraph("—",                               s1_val_num)],
            ]
            t = Table(rows_cal, colWidths=L_COL); t.setStyle(detail_ts)
            left.append(t)

            # ═══════════════════════════════════════════════════════════════════
            #  COLUMN 2 (RIGHT): Hasil Kalibrasi & Pelaksana
            # ═══════════════════════════════════════════════════════════════════
            right: List[Any] = []

            # ── 1. Hasil Kalibrasi ──
            right.append(Paragraph("Hasil Kalibrasi", s1_section))
            right.append(Paragraph("Calibration Results", s1_sub))

            verdict_color = Color(22/255, 163/255, 74/255) if self.result.overall_status == 'PASS' else Color(220/255, 38/255, 38/255)
            verdict_style = ParagraphStyle('Verdict', parent=styles['Normal'],
                fontName='Helvetica-Bold', fontSize=8.5, leading=11,
                textColor=verdict_color)
            rows_res = [
                [_lbl("Kesalahan Rentang Maks", "Max % Span Error"),      Paragraph(":", s1_colon), Paragraph(f"{self.result.max_error_pct_span:.4f}%", s1_val_num)],
                [_lbl("Akurasi Aktual", "Actual Accuracy"),               Paragraph(":", s1_colon), Paragraph(f"±{self.result.actual_accuracy_pct:.4f}%", s1_val_num)],
                [_lbl("Keputusan", "Overall Decision"),                   Paragraph(":", s1_colon), Paragraph(self.result.overall_status,                verdict_style)],
            ]
            t = Table(rows_res, colWidths=R_COL); t.setStyle(detail_ts)
            right.append(t)
            right.append(_hr())

            # ── 2. Pelaksana & Dokumen ──
            right.append(Paragraph("Pelaksana & Dokumen", s1_section))
            right.append(Paragraph("Personnel & Documentation", s1_sub))
            rows_doc = [
                [_lbl("Nama Operator", "Operator Name"),                   Paragraph(":", s1_colon), Paragraph(meta.operator_name,                    s1_val_text)],
                [_lbl("Halaman", "Page Number"),                          Paragraph(":", s1_colon), Paragraph("1 dari 2",                            s1_val_num)],
            ]
            t_doc = Table(rows_doc, colWidths=R_COL); t_doc.setStyle(detail_ts)
            right.append(t_doc)

            # ═══════════════════════════════════════════════════════════════════
            #  MASTER TABLE (two-column balanced layout)
            # ═══════════════════════════════════════════════════════════════════
            t_master = Table([[left, right]], colWidths=[244, 244])
            t_master.setStyle(TableStyle([
                ('VALIGN',        (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING',   (0, 0), (0, 0), 0),
                ('RIGHTPADDING',  (0, 0), (0, 0), 12),
                ('LEFTPADDING',   (1, 0), (1, 0), 12),
                ('RIGHTPADDING',  (1, 0), (1, 0), 0),
                ('TOPPADDING',    (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ]))
            elems.append(t_master)

            # ── Spacer to lower Signature Block to the bottom ──
            elems.append(Spacer(1, 140))

            # ── Signature Block (aligned right at bottom) ──
            sig_line = Table([['']], colWidths=[180])
            sig_line.setStyle(TableStyle([
                ('LINEBELOW', (0, 0), (-1, -1), 0.75, LABEL),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ]))

            sig_flowables = [
                sig_title,
                Spacer(1, 3),
                sig_loc_date,
                Spacer(1, 55),
                sig_line,
            ]

            sig_block_table = Table([[sig_flowables]], colWidths=[180])
            sig_block_table.setStyle(TableStyle([
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ]))

            sig_wrapper = Table([['', sig_block_table]], colWidths=[312, 180])
            sig_wrapper.setStyle(TableStyle([
                ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
                ('VALIGN', (0, 0), (-1, -1), 'BOTTOM'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ]))
            elems.append(sig_wrapper)
            elems.append(PageBreak())

        summary_data = [
            ["Parameter", "Value"],
            ["Span", f"{self.result.span:.4f}"],
            ["Max % Span Error", f"{self.result.max_error_pct_span:.4f}%"],
            ["Actual Accuracy", f"±{self.result.actual_accuracy_pct:.4f}%"],
            ["Overall Verdict", self.result.overall_status],
        ]
        t = Table(summary_data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), Color(44/255, 62/255, 80/255)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 24))

        point_header = ["Set Point", "Avg Reading", "Error %S", "Std Dev", "Repeat %S", "Status"]
        point_data = [point_header]
        for pr in self.result.point_results:
            point_data.append([
                f"{pr.set_point:.2f}", f"{pr.avg_reading:.4f}",
                f"{pr.error_pct_span:.4f}%", f"{pr.std_dev:.4f}",
                f"{pr.repeatability_pct_span:.4f}%", pr.status,
            ])
        t2 = Table(point_data)
        t2.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
        ]))
        elements.append(Paragraph("Detailed Calibration Points", styles['Heading2']))
        elements.append(t2)
        elements.append(Spacer(1, 24))
        elements.append(Paragraph("Performance Charts", styles['Heading2']))

        charts = CalibrationCharts(self.input_data, self.result)

        def get_chart_image(plot_method):
            fig, ax = plt.subplots(figsize=(6.5, 4.5))
            plot_method(ax)
            fig.tight_layout()
            buf = io.BytesIO()
            fig.savefig(buf, format='png', dpi=150)
            buf.seek(0)
            plt.close(fig)
            return Image(buf, width=450, height=311)

        elements.append(get_chart_image(charts.plot_calibration_curve))
        elements.append(Spacer(1, 12))
        elements.append(get_chart_image(charts.plot_error_curve))
        elements.append(Spacer(1, 12))
        elements.append(get_chart_image(charts.plot_residual_plot))
        doc.build(elements, canvasmaker=make_numbered_canvas(cert_number))

        buffer.seek(0)
        with open(filename, 'wb') as f:
            f.write(buffer.read())

    # ── Compare PDF ─────────────────────────────────────────────────────────
    def export_compare_pdf(self, filename: str, data: dict, cert_number: str = ''):
        import io
        import os
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.units import mm  # type: ignore
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            leftMargin=20*mm, rightMargin=20*mm,
            topMargin=39*mm, bottomMargin=45*mm
        )
        styles = getSampleStyleSheet()
        elements: List[Any] = []

        res_a  = data.get('resA', {})
        res_b  = data.get('resB', {})
        inp_a  = res_a.get('input', {})
        inp_b  = res_b.get('input', {})
        meta_a = inp_a.get('metadata') or {}
        meta_b = inp_b.get('metadata') or {}
        rA     = res_a.get('result', {})
        rB     = res_b.get('result', {})

        import datetime
        export_date = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # Spacer to prevent overlap with SLIDE1 certificate header
        elements.append(Spacer(1, 66))

        sub_style = ParagraphStyle('Sub', parent=styles['Normal'], alignment=1, textColor=colors.grey)
        elements.append(Paragraph(f"Comparison Report Generated: {export_date}", sub_style))
        elements.append(Spacer(1, 8))
        elements.append(Paragraph("Session Information", styles['Heading2']))

        sess_data = [
            ["", "Session A", "Session B"],
            ["Instrument", meta_a.get('instrument_model','—'), meta_b.get('instrument_model','—')],
            ["Operator",   meta_a.get('operator_name','—'),   meta_b.get('operator_name','—')],
            ["Cal. Date",  meta_a.get('date','—'),            meta_b.get('date','—')],
            ["Location",   meta_a.get('location','—'),        meta_b.get('location','—')],
            ["Status",     rA.get('overall_status','—'),      rB.get('overall_status','—')],
        ]
        t_sess = Table(sess_data, colWidths=[110, 200, 200])
        t_sess.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), Color(44/255, 62/255, 80/255)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        elements.append(t_sess)
        elements.append(Spacer(1, 16))
        elements.append(PageBreak())

        elements.append(Paragraph("Metric Comparison (Δ = B − A)", styles['Heading2']))

        def fmt(v, d=4):
            return f"{float(v):.{d}f}" if v is not None else "—"

        def delta_str(vA, vB):
            if vA is None or vB is None: return "—"
            d = float(vB) - float(vA)
            return f"{'+' if d > 0 else ''}{d:.5f}"

        metrics = [
            ["Metric", "Session A", "Session B", "Δ (B−A)"],
            ["Max Error (%FS)", fmt(rA.get('max_error_pct_span')), fmt(rB.get('max_error_pct_span')), delta_str(rA.get('max_error_pct_span'), rB.get('max_error_pct_span'))],
            ["Actual Accuracy (%)", fmt(rA.get('actual_accuracy_pct')), fmt(rB.get('actual_accuracy_pct')), delta_str(rA.get('actual_accuracy_pct'), rB.get('actual_accuracy_pct'))],
            ["Mean Bias Error", fmt(rA.get('mean_error'), 5), fmt(rB.get('mean_error'), 5), delta_str(rA.get('mean_error'), rB.get('mean_error'))],
            ["Std Deviation", fmt(rA.get('std_dev'), 5), fmt(rB.get('std_dev'), 5), delta_str(rA.get('std_dev'), rB.get('std_dev'))],
            ["RMSE", fmt(rA.get('rmse'), 5), fmt(rB.get('rmse'), 5), delta_str(rA.get('rmse'), rB.get('rmse'))],
            ["Linearity Error (%FS)", fmt(rA.get('linearity_error_pct')), fmt(rB.get('linearity_error_pct')), delta_str(rA.get('linearity_error_pct'), rB.get('linearity_error_pct'))],
            ["Zero Error", fmt(rA.get('zero_error'), 5), fmt(rB.get('zero_error'), 5), delta_str(rA.get('zero_error'), rB.get('zero_error'))],
            ["Span Error", fmt(rA.get('span_error'), 5), fmt(rB.get('span_error'), 5), delta_str(rA.get('span_error'), rB.get('span_error'))],
        ]
        t_metrics = Table(metrics, colWidths=[160, 100, 100, 100])
        t_metrics.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), Color(44/255, 62/255, 80/255)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f6fb')]),
        ]))
        elements.append(t_metrics)
        elements.append(Spacer(1, 16))

        elements.append(Paragraph("Point-by-Point Data", styles['Heading2']))
        ptsA = rA.get('point_results', [])
        ptsB = rB.get('point_results', [])
        max_len = max(len(ptsA), len(ptsB))
        pt_header = ["Set Point A", "Reading A", "Error A", "Set Point B", "Reading B", "Error B"]
        pt_data = [pt_header]
        for i in range(max_len):
            row: List[Any] = []
            if i < len(ptsA):
                pA = ptsA[i]
                row.extend([fmt(pA.get('set_point'), 2), fmt(pA.get('avg_reading'), 4), f"{fmt(pA.get('error_pct_span'), 4)}%"])
            else:
                row.extend(["—", "—", "—"])
            if i < len(ptsB):
                pB = ptsB[i]
                row.extend([fmt(pB.get('set_point'), 2), fmt(pB.get('avg_reading'), 4), f"{fmt(pB.get('error_pct_span'), 4)}%"])
            else:
                row.extend(["—", "—", "—"])
            pt_data.append(row)

        t_pts = Table(pt_data, colWidths=[75, 75, 75, 75, 75, 75])
        t_pts.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), Color(44/255, 62/255, 80/255)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f6fb')]),
        ]))
        elements.append(t_pts)
        elements.append(Spacer(1, 24))
        elements.append(Paragraph("Comparison Charts", styles['Heading2']))

        def get_comparison_chart(plot_type):
            fig, ax = plt.subplots(figsize=(6.5, 4.0))
            xi_A  = [p.get('set_point') for p in ptsA]
            yi_A  = [p.get('avg_reading') for p in ptsA]
            err_A = [p.get('error_pct_span') for p in ptsA]
            xi_B  = [p.get('set_point') for p in ptsB]
            yi_B  = [p.get('avg_reading') for p in ptsB]
            err_B = [p.get('error_pct_span') for p in ptsB]
            lrv   = float(inp_a.get('lrv', 0))
            urv   = float(inp_a.get('urv', 100))
            tol   = float(inp_a.get('tolerance_pct', 0.5))
            ax.grid(True, linestyle='-', linewidth=0.5, color='#dfe6e9', alpha=0.8)
            ax.spines['top'].set_visible(False); ax.spines['right'].set_visible(False)
            if plot_type == 'calibration':
                ax.set_title("Calibration Curve Comparison", fontsize=12, fontweight='bold', pad=10)
                ax.set_xlabel("Set Point Reference", fontsize=9)
                ax.set_ylabel("Measured Output", fontsize=9)
                if xi_A: ax.plot(xi_A, yi_A, marker='o', markersize=5, linestyle='-', color='#0984e3', label='Session A', alpha=0.8)
                if xi_B: ax.plot(xi_B, yi_B, marker='s', markersize=5, linestyle='-', color='#d63031', label='Session B', alpha=0.8)
                ax.plot([lrv, urv], [lrv, urv], color='#2d3436', linestyle='--', linewidth=1.5, label='Ideal', alpha=0.5)
            elif plot_type == 'error':
                ax.set_title("Error (% Span) Comparison", fontsize=12, fontweight='bold', pad=10)
                ax.set_xlabel("Set Point Reference", fontsize=9)
                ax.set_ylabel("Error (% of Span)", fontsize=9)
                if xi_A: ax.plot(xi_A, err_A, marker='o', markersize=5, linestyle='-', color='#0984e3', label='Session A Error')
                if xi_B: ax.plot(xi_B, err_B, marker='s', markersize=5, linestyle='-', color='#d63031', label='Session B Error')
                ax.axhline(y=tol,  color='#27ae60', linestyle='--', linewidth=1.2, label=f'Tolerance (±{tol}%)', alpha=0.6)
                ax.axhline(y=-tol, color='#27ae60', linestyle='--', linewidth=1.2, alpha=0.6)
                ax.axhline(y=0,    color='#2d3436', linewidth=1, alpha=0.4)
            ax.legend(frameon=True, fontsize=8)
            fig.tight_layout()
            buf = io.BytesIO()
            fig.savefig(buf, format='png', dpi=150)
            buf.seek(0)
            plt.close(fig)
            return Image(buf, width=450, height=277)

        elements.append(get_comparison_chart('calibration'))
        elements.append(Spacer(1, 12))
        elements.append(get_comparison_chart('error'))
        doc.build(elements)

        # Apply Template Background
        buffer.seek(0)
        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
        slide1_path = os.path.join(base_dir, 'template', 'SLIDE1.pdf')
        slidemore_path = os.path.join(base_dir, 'template', 'SLIDE_MORE.pdf')
        
        try:
            try:
                from pypdf import PdfReader, PdfWriter  # type: ignore
            except ImportError:
                from PyPDF2 import PdfReader, PdfWriter  # type: ignore
                
            writer: Any = PdfWriter()
            reader: Any = PdfReader(buffer)
            
            if os.path.exists(slide1_path) and os.path.exists(slidemore_path):
                for idx, page in enumerate(reader.pages):
                    t_path = slide1_path if idx == 0 else slidemore_path
                    temp_reader: Any = PdfReader(t_path)
                    bg_page: Any = temp_reader.pages[0]
                    bg_page.merge_page(page)
                    writer.add_page(bg_page)
            else:
                for page in reader.pages:
                    writer.add_page(page)
                    
            with open(filename, 'wb') as f:
                writer.write(f)
                
        except ImportError:
            buffer.seek(0)
            with open(filename, 'wb') as f:
                f.write(buffer.read())
